"""Synthetic test set: geometric and photometric variants of the open-set videos.

Each variant keeps duration, fps and frame timing, so the original labels (verdict + timestamp) stay valid.

usage: python tools/augment.py SRC_DIR OUT_DIR [--variants mirror,persp_left,...] [--workers 6]
SRC_DIR must contain videos/ and open_labels.xlsx.
"""
import argparse
import os
import re
from multiprocessing import Pool

import cv2
import numpy as np
import openpyxl


def _persp(img, shift_x, shift_y):
    """Warp as if the camera moved sideways/up: top edge narrower on one side."""
    h, w = img.shape[:2]
    src = np.float32([[0, 0], [w, 0], [w, h], [0, h]])
    dx, dy = shift_x * w, shift_y * h
    if shift_x >= 0:
        dst = np.float32([[dx, dy], [w, 0], [w, h], [0, h]])
    else:
        dst = np.float32([[0, 0], [w + dx, dy], [w, h], [0, h]])
    m = cv2.getPerspectiveTransform(src, dst)
    out = cv2.warpPerspective(img, m, (w, h), borderMode=cv2.BORDER_REFLECT)
    # zoom in a little so reflected borders are mostly cropped away
    return _zoom(out, 0.8)


def _zoom(img, keep):
    h, w = img.shape[:2]
    ch, cw = int(h * keep), int(w * keep)
    y0, x0 = (h - ch) // 2, (w - cw) // 2
    return cv2.resize(img[y0:y0 + ch, x0:x0 + cw], (w, h), interpolation=cv2.INTER_LINEAR)


def _rotate(img, deg):
    h, w = img.shape[:2]
    m = cv2.getRotationMatrix2D((w / 2, h / 2), deg, 1.18)
    return cv2.warpAffine(img, m, (w, h), borderMode=cv2.BORDER_REFLECT)


def _gamma(img, g):
    lut = np.array([((i / 255.0) ** g) * 255 for i in range(256)], np.uint8)
    return cv2.LUT(img, lut)


def _dark(img):
    out = _gamma(img, 1.7).astype(np.float32)
    out[..., 0] *= 1.08  # slight blue cast of an evening/cheap camera
    out[..., 2] *= 0.92
    return np.clip(out, 0, 255).astype(np.uint8)


def _bright(img):
    return cv2.convertScaleAbs(img, alpha=1.3, beta=18)


def _lowq(img, rng):
    h, w = img.shape[:2]
    small = cv2.resize(img, (w // 2, h // 2), interpolation=cv2.INTER_AREA)
    out = cv2.resize(small, (w, h), interpolation=cv2.INTER_LINEAR)
    out = cv2.GaussianBlur(out, (3, 3), 0).astype(np.float32)
    out += rng.normal(0, 8, out.shape)
    return np.clip(out, 0, 255).astype(np.uint8)


VARIANTS = {
    "mirror": lambda img, rng: cv2.flip(img, 1),
    "persp_left": lambda img, rng: _persp(img, 0.14, 0.06),
    "persp_right": lambda img, rng: _persp(img, -0.14, 0.06),
    "rotate": lambda img, rng: _rotate(img, 8),
    "zoom": lambda img, rng: _zoom(img, 0.8),
    "dark": lambda img, rng: _dark(img),
    "bright": lambda img, rng: _bright(img),
    "lowq": lambda img, rng: _lowq(img, rng),
    # hardest: other side of the room + other angle + bad light
    "combo": lambda img, rng: _dark(_persp(cv2.flip(img, 1), 0.12, 0.05)),
}


def process(job):
    src, dst, variant = job
    if os.path.exists(dst):
        return dst
    cap = cv2.VideoCapture(src)
    fps = cap.get(cv2.CAP_PROP_FPS) or 30.0
    w, h = int(cap.get(cv2.CAP_PROP_FRAME_WIDTH)), int(cap.get(cv2.CAP_PROP_FRAME_HEIGHT))
    tmp = dst + ".part.mov"
    out = cv2.VideoWriter(tmp, cv2.VideoWriter_fourcc(*"mp4v"), fps, (w, h))
    rng = np.random.default_rng(abs(hash((os.path.basename(src), variant))) % 2**32)
    fn = VARIANTS[variant]
    while True:
        ok, img = cap.read()
        if not ok:
            break
        out.write(fn(img, rng))
    cap.release()
    out.release()
    os.replace(tmp, dst)
    return dst


def norm(s):
    return s.strip().replace("Х", "X")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("src")
    ap.add_argument("out")
    ap.add_argument("--variants", default=",".join(VARIANTS))
    ap.add_argument("--workers", type=int, default=max(1, (os.cpu_count() or 4) - 2))
    a = ap.parse_args()
    variants = a.variants.split(",")

    wb = openpyxl.load_workbook(os.path.join(a.src, "open_labels.xlsx"))
    rows = [r for r in list(wb.active.iter_rows(values_only=True))[1:] if r[0]]
    files = {norm(f[:-4]): f for f in os.listdir(os.path.join(a.src, "videos")) if f.endswith(".mov")}
    # byte-identical duplicates in the open set are generated once
    skip = {"X-47", "X-48"}

    jobs, label_rows = [], []
    for name, ts, lab in rows:
        n = norm(name)
        if n in skip or n not in files:
            continue
        for v in variants:
            os.makedirs(os.path.join(a.out, v), exist_ok=True)
            dst = os.path.join(a.out, v, f"{n}_{v}.mov")
            jobs.append((os.path.join(a.src, "videos", files[n]), dst, v))
            label_rows.append((f"{n}_{v}", v, n, ts, lab))

    with Pool(a.workers) as p:
        for i, d in enumerate(p.imap_unordered(process, jobs), 1):
            if i % 20 == 0 or i == len(jobs):
                print(f"{i}/{len(jobs)} {os.path.basename(d)}", flush=True)

    out = openpyxl.Workbook()
    ws = out.active
    ws.title = "labels"
    ws.append(["наименование файла", "вариант", "исходное видео", "таймстемп", "нарушение / без нарушения"])
    for r in label_rows:
        ws.append(list(r))
    out.save(os.path.join(a.out, "synthetic_labels.xlsx"))
    print("labels:", len(label_rows))


if __name__ == "__main__":
    main()
