"""Evaluate a deployed service through the jury API on the synthetic (and original) videos.

usage:
  export CASHGUARD_API_KEY=...            # never pass the key on the command line
  python tools/api_eval.py https://your-host --data ~/open_synthetic \
         --variants original,combo,persp_left,dark [--per-variant 10] [--out api_results.json]

"original" uses ~/open/videos with open_labels.xlsx. Results are saved after every video,
so the run can be stopped and resumed (already finished videos are skipped).
"""
import argparse
import json
import mimetypes
import os
import re
import time
import urllib.request
import uuid

import openpyxl

TIMEOUT_SEC = 300


def norm(s):
    return s.strip().replace("Х", "X")


def parse_ts(s):
    if not s:
        return None
    s = str(s)
    m = re.search(r"(\d+)\s*мин\w*\s*(\d+)", s)
    if m:
        return int(m.group(1)) * 60 + int(m.group(2))
    return int(re.search(r"\d+", s).group())


def post_video(base, key, path):
    boundary = uuid.uuid4().hex
    with open(path, "rb") as f:
        data = f.read()
    body = (f"--{boundary}\r\nContent-Disposition: form-data; name=\"file\"; filename=\"{os.path.basename(path)}\"\r\n"
            f"Content-Type: {mimetypes.guess_type(path)[0] or 'video/quicktime'}\r\n\r\n").encode() + data + \
           f"\r\n--{boundary}--\r\n".encode()
    req = urllib.request.Request(base.rstrip("/") + "/api/v1/analyze", data=body, method="POST",
                                 headers={"X-API-Key": key, "Content-Type": f"multipart/form-data; boundary={boundary}"})
    with urllib.request.urlopen(req, timeout=120) as r:
        return json.load(r)["job_id"]


def get_job(base, key, job_id):
    req = urllib.request.Request(base.rstrip("/") + f"/api/v1/jobs/{job_id}", headers={"X-API-Key": key})
    with urllib.request.urlopen(req, timeout=60) as r:
        return json.load(r)


def run_one(base, key, path):
    t0 = time.time()
    try:
        job = post_video(base, key, path)
    except Exception as e:  # noqa: BLE001
        return {"status": "failed", "error": f"POST: {e}", "sec": round(time.time() - t0, 1)}
    while time.time() - t0 < TIMEOUT_SEC:
        time.sleep(3)
        try:
            res = get_job(base, key, job)
        except Exception as e:  # noqa: BLE001
            res = {"status": "processing", "poll_error": str(e)}
        if res.get("status") in ("done", "failed"):
            res["sec"] = round(time.time() - t0, 1)
            res.pop("violations_frames", None)
            for v in res.get("violations") or []:
                v.pop("evidence_frames", None)  # keep the results file small
            return res
    return {"status": "timeout", "sec": round(time.time() - t0, 1)}


def load_items(data_dir, variants, per_variant):
    items = []
    if "original" in variants:
        wb = openpyxl.load_workbook(os.path.expanduser("~/open/open_labels.xlsx"))
        vids = os.path.expanduser("~/open/videos")
        files = {norm(f[:-4]): os.path.join(vids, f) for f in os.listdir(vids) if f.endswith(".mov")}
        rows = [r for r in list(wb.active.iter_rows(values_only=True))[1:] if r[0] and norm(r[0]) not in ("X-47", "X-48")]
        for name, ts, lab in rows[:per_variant or None]:
            items.append(("original", norm(name), files[norm(name)], lab.strip() == "нарушение", parse_ts(ts)))
    wb = openpyxl.load_workbook(os.path.join(data_dir, "synthetic_labels.xlsx"))
    per = {}
    for fname, variant, src, ts, lab in list(wb.active.iter_rows(values_only=True))[1:]:
        if variant not in variants:
            continue
        per.setdefault(variant, []).append(
            (variant, src, os.path.join(data_dir, variant, fname + ".mov"), lab.strip() == "нарушение", parse_ts(ts)))
    for v, lst in per.items():
        items += lst[:per_variant or None]
    return items


def report(results):
    by = {}
    for r in results.values():
        by.setdefault(r["variant"], []).append(r)
    print(f"{'variant':12} {'n':>3} {'acc':>5} {'TP':>3} {'FN':>3} {'FP':>3} {'TN':>3} {'ts<=2s':>6} {'fail':>4} {'max_s':>6}")
    for v, rows in sorted(by.items()):
        tp = fn = fp = tn = fail = ts_ok = 0
        for r in rows:
            res = r["result"]
            if res.get("status") != "done":
                fail += 1
            pred = res.get("verdict") == "violation"
            if r["violation"] and pred:
                tp += 1
                tcs = [x.get("timecode_sec") for x in res.get("violations") or [] if x.get("timecode_sec") is not None]
                if r["ts"] is not None and tcs and min(abs(t - r["ts"]) for t in tcs) <= 2:
                    ts_ok += 1
            elif r["violation"]:
                fn += 1
            elif pred:
                fp += 1
            else:
                tn += 1
        n = len(rows)
        print(f"{v:12} {n:3} {(tp + tn) / n:5.2f} {tp:3} {fn:3} {fp:3} {tn:3} {ts_ok:6} {fail:4} "
              f"{max(r['result'].get('sec', 0) for r in rows):6.0f}")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("base_url")
    ap.add_argument("--data", default=os.path.expanduser("~/open_synthetic"))
    ap.add_argument("--variants", default="original,combo,persp_left,dark")
    ap.add_argument("--per-variant", type=int, default=0, help="limit videos per variant (0 = all)")
    ap.add_argument("--out", default="api_results.json")
    a = ap.parse_args()
    key = os.environ.get("CASHGUARD_API_KEY")
    if not key:
        raise SystemExit("set CASHGUARD_API_KEY first")

    results = json.load(open(a.out)) if os.path.exists(a.out) else {}
    items = load_items(a.data, set(a.variants.split(",")), a.per_variant)
    for i, (variant, name, path, viol, ts) in enumerate(items, 1):
        k = f"{variant}/{name}"
        if k in results:
            continue
        res = run_one(a.base_url, key, path)
        results[k] = {"variant": variant, "name": name, "violation": viol, "ts": ts, "result": res}
        print(f"[{i}/{len(items)}] {k} gt={'V' if viol else '-'}@{ts} -> {res.get('status')} "
              f"{res.get('verdict')} {[x.get('timecode_sec') for x in res.get('violations') or []]} ({res.get('sec')}s)",
              flush=True)
        json.dump(results, open(a.out, "w"), ensure_ascii=False, indent=1)
    report(results)


if __name__ == "__main__":
    main()
